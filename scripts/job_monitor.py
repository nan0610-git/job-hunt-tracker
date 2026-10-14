#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
岗位监控过滤与清单管理 (job-hunt-tracker)

子命令:
  config    查看 / 初始化 / 修改岗位监控配置
  filter    对浏览器抓取的原始岗位列表做 过滤+去重，生成待投清单
  seen      将待投清单中的岗位标记为已处理（不再重复推荐）
  show      查看当前待投清单
  clear     清空待投清单

原始岗位 JSON 格式（由浏览器抓取步骤产出，抓取要点见 references/job-monitor-guide.md）:
[
  {"platform": "zhipin", "title": "运维工程师", "company": "云启科技",
   "salary": "15-25K·14薪", "city": "西安", "url": "https://...",
   "tags": ["3-5年", "本科"], "published": "2026-10-14"}
]
"""

import argparse
import json
import os
import re
import sys
from datetime import datetime

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

DEFAULT_CONFIG = {
    "keywords": ["运维工程师", "SRE", "系统运维"],
    "city": "西安",
    "salary_min_k": 15,
    "exclude_companies": [],
    "exclude_keywords": ["外包", "驻场", "劳务", "派遣"],
    "daily_limit": 30,
    "platforms": ["zhipin", "zhilian"]
}

CONFIG_FILE = "config.json"
SEEN_FILE = "seen_jobs.json"
TODO_FILE = "todo_jobs.json"


def data_path(data_dir, name):
    return os.path.join(data_dir, name)


def load_json(path, default):
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except (json.JSONDecodeError, OSError):
            print(f"⚠️ 配置/数据文件损坏，已重置: {path}")
    return default


def save_json(path, data):
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def parse_salary_min(salary):
    """解析薪资字符串，返回最低 K 数；无法解析返回 None。"""
    if not salary:
        return None
    s = str(salary).replace(",", "").upper()
    m = re.search(r"(\d+(?:\.\d+)?)\s*[-~]\s*(\d+(?:\.\d+)?)\s*K", s)
    if m:
        return float(m.group(1))
    m = re.search(r"(\d+(?:\.\d+)?)\s*K", s)
    if m:
        return float(m.group(1))
    m = re.search(r"(\d+)\s*万\s*/\s*年", s)
    if m:
        return round(float(m.group(1)) * 10 / 12, 1)
    return None


def job_key(job):
    return f"{str(job.get('company', '')).strip()}|{str(job.get('title', '')).strip()}"


def load_tracker_jobs(tracker):
    """读跟踪表已投递记录，返回 (company|title) 集合。"""
    applied = set()
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("⚠️ 未安装 openpyxl，跳过与跟踪表的去重对比（pip install openpyxl）")
        return applied
    try:
        wb = load_workbook(tracker, read_only=True)
        ws = wb["投递记录"]
        for row in range(2, ws.max_row + 1):
            company = ws.cell(row=row, column=1).value
            if company is None:
                continue
            title = str(ws.cell(row=row, column=2).value or "").strip()
            applied.add(f"{str(company).strip()}|{title}")
    except Exception as e:  # noqa: BLE001
        print(f"⚠️ 读取跟踪表失败，跳过其去重对比: {e}")
    return applied


def cmd_config(args):
    path = data_path(args.data_dir, CONFIG_FILE)
    cfg = load_json(path, None)
    if args.init:
        if cfg:
            print(f"配置已存在: {path}")
        else:
            cfg = dict(DEFAULT_CONFIG)
            save_json(path, cfg)
            print(f"已生成默认配置: {path}")
    if args.set:
        if not cfg:
            cfg = dict(DEFAULT_CONFIG)
        for kv in args.set:
            if "=" not in kv:
                print(f"--set 格式应为 key=value（JSON 值）: {kv}")
                sys.exit(1)
            key, _, val = kv.partition("=")
            try:
                cfg[key.strip()] = json.loads(val)
            except json.JSONDecodeError:
                cfg[key.strip()] = val
        save_json(path, cfg)
    if not cfg:
        print(f"尚未初始化配置，运行: {os.path.basename(sys.argv[0])} config --init")
        sys.exit(1)
    print(json.dumps(cfg, ensure_ascii=False, indent=2))


def cmd_filter(args):
    cfg = load_json(data_path(args.data_dir, CONFIG_FILE), None)
    if not cfg:
        print("尚未初始化配置，先运行 config --init")
        sys.exit(1)
    raw = load_json(args.raw, None)
    if raw is None:
        print(f"无法读取原始岗位文件: {args.raw}")
        sys.exit(1)

    seen_keys = set(load_json(data_path(args.data_dir, SEEN_FILE), []))
    todo = load_json(data_path(args.data_dir, TODO_FILE), [])
    todo_keys = {job_key(j) for j in todo}
    applied = load_tracker_jobs(args.tracker) if args.tracker else set()

    kept, skipped = [], {"keyword": 0, "salary": 0, "company": 0, "exclude_kw": 0,
                         "seen": 0, "in_todo": 0, "applied": 0}
    for job in raw:
        title = str(job.get("title", "")).strip()
        company = str(job.get("company", "")).strip()
        if not title or not company:
            continue
        if not any(kw.lower() in title.lower() for kw in cfg["keywords"]):
            skipped["keyword"] += 1
            continue
        sal_min = parse_salary_min(job.get("salary", ""))
        job["_salary_min_k"] = sal_min
        if sal_min is not None and sal_min < cfg["salary_min_k"]:
            skipped["salary"] += 1
            continue
        if any(ex and ex in company for ex in cfg["exclude_companies"]):
            skipped["company"] += 1
            continue
        blob = title + " " + " ".join(str(t) for t in job.get("tags", []))
        if any(ex and ex in blob for ex in cfg["exclude_keywords"]):
            skipped["exclude_kw"] += 1
            continue
        key = job_key(job)
        if key in seen_keys:
            skipped["seen"] += 1
            continue
        if key in todo_keys:
            skipped["in_todo"] += 1
            continue
        if key in applied:
            skipped["applied"] += 1
            continue
        kept.append(job)

    kept.sort(key=lambda j: -(j.get("_salary_min_k") or 0))
    limit = cfg.get("daily_limit", 30)
    if limit > 0:
        kept = kept[:limit]

    todo.extend(kept)
    save_json(data_path(args.data_dir, TODO_FILE), todo)

    print(f"原始 {len(raw)} 条 → 保留 {len(kept)} 条"
          f"（过滤: 关键词 {skipped['keyword']} / 薪资 {skipped['salary']}"
          f" / 黑名单公司 {skipped['company']} / 排除词 {skipped['exclude_kw']}"
          f" / 已见过 {skipped['seen']} / 已在清单 {skipped['in_todo']}"
          f" / 已投递过 {skipped['applied']}）\n")
    if kept:
        print("待投清单（按薪资降序）:")
        for i, j in enumerate(kept, 1):
            note = "" if j.get("_salary_min_k") is not None else " [面议]"
            print(f"{i}. [{j.get('platform', '?')}] {j.get('company')} - {j.get('title')}"
                  f" | {j.get('salary', '面议')}{note} | {j.get('city', '')}")
            print(f"   {j.get('url', '')}")
    else:
        print("本轮没有新的符合条件岗位。")


def cmd_seen(args):
    path = data_path(args.data_dir, TODO_FILE)
    todo = load_json(path, [])
    seen_path = data_path(args.data_dir, SEEN_FILE)
    seen = load_json(seen_path, [])
    seen_keys = set(seen)
    if args.all:
        marked = todo
        todo = []
    elif args.url:
        marked = [j for j in todo if str(j.get("url", "")) == args.url]
        todo = [j for j in todo if str(j.get("url", "")) != args.url]
    else:
        print("用 --all 标记全部，或 --url <岗位链接> 标记单条")
        sys.exit(1)
    for j in marked:
        k = job_key(j)
        if k not in seen_keys:
            seen.append(k)
            seen_keys.add(k)
    save_json(seen_path, seen)
    save_json(path, todo)
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    print(f"已标记 {len(marked)} 条为已处理（{stamp}），待投清单剩余 {len(todo)} 条。")


def cmd_show(args):
    todo = load_json(data_path(args.data_dir, TODO_FILE), [])
    if not todo:
        print("待投清单为空。")
        return
    print(f"待投清单共 {len(todo)} 条:")
    for i, j in enumerate(todo, 1):
        print(f"{i}. [{j.get('platform', '?')}] {j.get('company')} - {j.get('title')}"
              f" | {j.get('salary', '面议')} | {j.get('city', '')}")
        print(f"   {j.get('url', '')}")


def cmd_clear(args):
    path = data_path(args.data_dir, TODO_FILE)
    old = load_json(path, [])
    save_json(path, [])
    print(f"已清空待投清单（原 {len(old)} 条）。")


def main():
    parser = argparse.ArgumentParser(description="岗位监控过滤与清单管理")
    parser.add_argument("--data-dir", default=".job-monitor", help="数据目录（默认 ./.job-monitor）")
    sub = parser.add_subparsers(dest="command", required=True)

    p = sub.add_parser("config", help="查看/初始化/修改配置")
    p.add_argument("--init", action="store_true", help="初始化默认配置")
    p.add_argument("--set", nargs="*", default=None, help="修改配置项，如 --set salary_min_k=18 'keywords=[\"运维\"]'")
    p.set_defaults(func=cmd_config)

    p = sub.add_parser("filter", help="过滤+去重，生成待投清单")
    p.add_argument("raw", help="浏览器抓取的原始岗位 JSON 文件")
    p.add_argument("--tracker", default=None, help="求职跟踪表 xlsx 路径（用于已投递去重）")
    p.set_defaults(func=cmd_filter)

    p = sub.add_parser("seen", help="标记岗位为已处理")
    p.add_argument("--all", action="store_true")
    p.add_argument("--url", default=None)
    p.set_defaults(func=cmd_seen)

    p = sub.add_parser("show", help="查看待投清单")
    p.set_defaults(func=cmd_show)

    p = sub.add_parser("clear", help="清空待投清单")
    p.set_defaults(func=cmd_clear)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
