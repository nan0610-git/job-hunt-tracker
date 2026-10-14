#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
求职投递跟踪表管理脚本 (job-hunt-tracker)

子命令:
  init        初始化跟踪表（工作表：投递记录 / Offer对比）
  add         添加一条投递记录
  update      更新某条投递的状态/备注
  list        列出投递记录（可按状态过滤）
  stats       投递统计概览
  add-offer   录入一条 Offer 数据
  compare     Offer 加权对比

依赖: openpyxl (pip install openpyxl)
"""

import argparse
import os
import sys
from datetime import date, datetime, timedelta

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    print("缺少依赖 openpyxl，请先安装: pip install openpyxl")
    sys.exit(1)

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

DEFAULT_FILE = "求职跟踪表.xlsx"
SHEET_APP = "投递记录"
SHEET_OFFER = "Offer对比"

STATUS_OPTIONS = ["投递中", "笔试", "一面", "二面", "HR面", "Offer", "已拒绝", "已婉拒", "无回音"]
TERMINAL_STATUS = {"Offer", "已拒绝", "已婉拒"}

APP_HEADERS = ["公司", "岗位", "工作地点", "投递渠道", "投递日期", "当前状态",
               "状态更新时间", "薪资范围", "下次跟进日期", "备注"]
OFFER_HEADERS = ["公司", "岗位", "月薪(元)", "年薪(万)", "股票期权(万/年)", "年终奖(月)",
                 "五险一金比例", "技术匹配(1-5)", "成长空间(1-5)", "稳定性(1-5)",
                 "工作轻松度(1-5)", "通勤便利(1-5)", "加权总分", "备注"]
# --scores 顺序: 技术匹配, 成长空间, 稳定性, 工作轻松度, 通勤便利
SCORE_COLS = [8, 9, 10, 11, 12]  # 对应 OFFER_HEADERS 的列号(1-based)

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF")

TODAY = date.today()


def parse_date(s):
    """解析日期字符串，支持 yyyy-mm-dd / yyyy/m/d。"""
    if not s:
        return None
    s = s.strip().replace("/", "-")
    for fmt in ("%Y-%m-%d", "%Y-%m", "%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"无法解析日期: {s}（支持格式如 2026-10-14）")


def fmt_date(d):
    return d.strftime("%Y-%m-%d") if isinstance(d, (date, datetime)) else (d or "")


def style_sheet(ws, headers, status_col_letter=None):
    ws.append(headers)
    for idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = 16
    ws.column_dimensions[ws.cell(row=1, column=len(headers)).column_letter].width = 30
    ws.freeze_panes = "A2"
    if status_col_letter:
        dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_OPTIONS) + '"',
                            allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{status_col_letter}2:{status_col_letter}500")


def next_empty_row(ws):
    row = ws.max_row + 1
    if ws.max_row == 1 and ws.cell(row=1, column=1).value is None:
        row = 2
    return row


def find_rows(ws, company, role=None):
    found = []
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row=row, column=1).value or "").strip() == company:
            if role and str(ws.cell(row=row, column=2).value or "").strip() != role:
                continue
            found.append(row)
    return found


def to_float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------- 子命令实现

def cmd_init(args):
    path = args.file
    if os.path.exists(path) and not args.force:
        print(f"文件已存在: {path}（如需重建请加 --force，注意旧数据将丢失）")
        sys.exit(1)
    wb = Workbook()
    ws1 = wb.active
    ws1.title = SHEET_APP
    style_sheet(ws1, APP_HEADERS, status_col_letter="F")
    ws2 = wb.create_sheet(SHEET_OFFER)
    style_sheet(ws2, OFFER_HEADERS)
    wb.save(path)
    print(f"已初始化跟踪表: {path}")
    print(f"  - {SHEET_APP}: 记录每次投递与进度")
    print(f"  - {SHEET_OFFER}: 录入 Offer 用于对比决策")


def cmd_add(args):
    d = parse_date(args.date) or TODAY
    followup = parse_date(args.followup) or (d + timedelta(days=7))
    wb = load_workbook(args.file)
    ws = wb[SHEET_APP]
    row = next_empty_row(ws)
    values = [args.company, args.role or "", args.location or "", args.channel or "",
              d, args.status or "投递中", datetime.now(), args.salary or "",
              followup, args.note or ""]
    for col, val in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=val)
        if isinstance(val, (date, datetime)):
            cell.number_format = "yyyy-mm-dd"
    wb.save(args.file)
    print(f"已添加投递记录: {args.company} - {args.role or '未填写岗位'}（状态: {args.status or '投递中'}）")
    print(f"建议跟进日期: {fmt_date(followup)}")


def cmd_update(args):
    wb = load_workbook(args.file)
    ws = wb[SHEET_APP]
    rows = find_rows(ws, args.company, args.role)
    if not rows:
        print(f"未找到公司为「{args.company}」"
              + (f" 且岗位为「{args.role}」" if args.role else "")
              + " 的记录。现有记录：")
        for r in range(2, ws.max_row + 1):
            print(f"  - {ws.cell(row=r, column=1).value} | {ws.cell(row=r, column=2).value}")
        sys.exit(1)
    if len(rows) > 1 and not args.role:
        print("该公司有多条记录，请用 --role 指定岗位：")
        for r in rows:
            print(f"  - {ws.cell(row=r, column=2).value}")
        sys.exit(1)
    row = rows[0]
    if args.status:
        if args.status not in STATUS_OPTIONS:
            print(f"无效状态「{args.status}」，可选: {' / '.join(STATUS_OPTIONS)}")
            sys.exit(1)
        ws.cell(row=row, column=6, value=args.status)
        ws.cell(row=row, column=7, value=datetime.now()).number_format = "yyyy-mm-dd"
    if args.note:
        old = str(ws.cell(row=row, column=10).value or "").strip()
        stamp = TODAY.strftime("%Y-%m-%d")
        add = f"【{stamp}】{args.note}"
        ws.cell(row=row, column=10, value=(old + "\n" + add).strip())
    wb.save(args.file)
    print(f"已更新: {ws.cell(row=row, column=1).value} - {ws.cell(row=row, column=2).value}"
          + (f" → {args.status}" if args.status else ""))


def cmd_list(args):
    wb = load_workbook(args.file, read_only=True)
    ws = wb[SHEET_APP]
    count = 0
    for row in range(2, ws.max_row + 1):
        company = ws.cell(row=row, column=1).value
        if company is None:
            continue
        status = str(ws.cell(row=row, column=6).value or "").strip()
        if args.status and status != args.status:
            continue
        count += 1
        print(" | ".join(str(x) for x in [
            company,
            ws.cell(row=row, column=2).value or "",
            ws.cell(row=row, column=4).value or "",
            fmt_date(ws.cell(row=row, column=5).value),
            status,
            ws.cell(row=row, column=9).value if ws.cell(row=row, column=9).value else "",
        ]))
    print(f"共 {count} 条记录" + (f"（状态={args.status}）" if args.status else ""))


def cmd_stats(args):
    wb = load_workbook(args.file, read_only=True)
    ws = wb[SHEET_APP]
    total = 0
    status_count = {}
    week_new = 0
    followup_due = []
    for row in range(2, ws.max_row + 1):
        company = ws.cell(row=row, column=1).value
        if company is None:
            continue
        total += 1
        status = str(ws.cell(row=row, column=6).value or "").strip() or "未填写"
        status_count[status] = status_count.get(status, 0) + 1
        d = ws.cell(row=row, column=5).value
        if isinstance(d, datetime):
            d = d.date()
        if isinstance(d, date) and (TODAY - d).days <= 7:
            week_new += 1
        fu = ws.cell(row=row, column=9).value
        if isinstance(fu, datetime):
            fu = fu.date()
        if isinstance(fu, date) and fu <= TODAY and status not in TERMINAL_STATUS:
            followup_due.append(f"{company} - {ws.cell(row=row, column=2).value or ''}")
    print(f"总投递数: {total}")
    print(f"本周新增: {week_new}")
    print("各状态:")
    for s in STATUS_OPTIONS:
        if status_count.get(s):
            print(f"  {s}: {status_count[s]}")
    for s, c in status_count.items():
        if s not in STATUS_OPTIONS:
            print(f"  {s}: {c}")
    if status_count.get("无回音") and total:
        ratio = status_count["无回音"] / total
        if ratio > 0.4:
            print(f"⚠️ 「无回音」占比 {ratio:.0%}，建议优化简历或拓宽投递渠道")
    if followup_due:
        print("待跟进:")
        for item in followup_due:
            print(f"  - {item}")


def cmd_add_offer(args):
    wb = load_workbook(args.file)
    ws = wb[SHEET_OFFER]
    scores = []
    if args.scores:
        parts = [p.strip() for p in args.scores.split(",")]
        if len(parts) != 5:
            print("--scores 需要恰好 5 个分数（技术匹配,成长空间,稳定性,工作轻松度,通勤便利）")
            sys.exit(1)
        for p in parts:
            v = float(p)
            if not 1 <= v <= 5:
                print(f"分数 {v} 超出范围 1-5")
                sys.exit(1)
            scores = [to_float(p) for p in parts]
    annual = args.annual
    if annual is None and args.month_salary:
        bonus_months = args.bonus_months or 0
        annual = round(args.month_salary * (12 + bonus_months) / 10000, 2)
    row = next_empty_row(ws)
    values = [args.company, args.role or "", args.month_salary, annual, args.stock,
              args.bonus_months, args.social_insurance or ""]
    if scores:
        values += scores
    values += [None, args.note or ""]
    values += [None] * (len(OFFER_HEADERS) - len(values)) if len(values) < len(OFFER_HEADERS) else []
    for col, val in enumerate(values[:len(OFFER_HEADERS)], start=1):
        if val is not None:
            ws.cell(row=row, column=col, value=val)
    wb.save(args.file)
    print(f"已录入 Offer: {args.company} - {args.role or '未填写岗位'}")
    if annual:
        print(f"折算年薪: {annual} 万（如需含股票请直接在表中补充）")


def cmd_compare(args):
    weights = [float(x) for x in args.weights.split(",")]
    if len(weights) != 3 or abs(sum(weights) - 1.0) > 1e-6:
        print("--weights 需要 3 个权重且总和为 1，如 0.4,0.3,0.3（薪资包,发展,平衡）")
        sys.exit(1)
    wb = load_workbook(args.file, read_only=True)
    ws = wb[SHEET_OFFER]
    offers = []
    for row in range(2, ws.max_row + 1):
        company = ws.cell(row=row, column=1).value
        if company is None:
            continue
        vals = {col: to_float(ws.cell(row=row, column=col).value) for col in SCORE_COLS}
        annual = to_float(ws.cell(row=row, column=4).value)
        monthly = to_float(ws.cell(row=row, column=3).value)
        bonus = to_float(ws.cell(row=row, column=6).value) or 0
        if annual is None and monthly:
            annual = round(monthly * (12 + bonus) / 10000, 2)
        offers.append({
            "company": str(company).strip(),
            "role": str(ws.cell(row=row, column=2).value or "").strip(),
            "annual": annual,
            "scores": vals,
            "row": row,
        })
    if len(offers) < 2:
        print(f"Offer 对比至少需要 2 条记录（当前 {len(offers)} 条），请先用 add-offer 录入。")
        sys.exit(1)

    missing = [o["company"] for o in offers if o["annual"] is None]
    if missing:
        print(f"⚠️ 以下 Offer 缺少年薪数据，薪资分按 0 处理: {', '.join(missing)}")
    max_annual = max((o["annual"] or 0) for o in offers) or 1
    dim_names = {8: "技术匹配", 9: "成长空间", 10: "稳定性", 11: "工作轻松度", 12: "通勤便利"}

    results = []
    for o in offers:
        salary_score = (o["annual"] or 0) / max_annual * 5
        dev_vals = [o["scores"][8], o["scores"][9]]
        bal_vals = [o["scores"][10], o["scores"][11], o["scores"][12]]
        dev_score = sum(v for v in dev_vals if v is not None) / max(1, len([v for v in dev_vals if v is not None]))
        bal_score = sum(v for v in bal_vals if v is not None) / max(1, len([v for v in bal_vals if v is not None]))
        total = weights[0] * salary_score + weights[1] * dev_score + weights[2] * bal_score
        # 最强/最弱项（仅看人工评分维度）
        rated = {dim_names[c]: v for c, v in o["scores"].items() if v is not None}
        best = max(rated, key=rated.get) if rated else "—"
        worst = min(rated, key=rated.get) if rated else "—"
        results.append({
            "company": o["company"], "role": o["role"], "annual": o["annual"],
            "salary_score": round(salary_score, 2),
            "dev_score": round(dev_score, 2), "bal_score": round(bal_score, 2),
            "total": round(total, 2), "best": best, "worst": worst,
        })

    results.sort(key=lambda x: -x["total"])
    print(f"权重: 薪资包 {weights[0]:.0%} / 发展前景 {weights[1]:.0%} / 工作生活平衡 {weights[2]:.0%}\n")
    header = f"{'排名':<4}{'公司':<14}{'岗位':<12}{'年薪(万)':<10}{'薪资分':<8}{'发展分':<8}{'平衡分':<8}{'总分':<8}"
    print(header)
    print("-" * len(header))
    for i, r in enumerate(results, 1):
        annual = r["annual"] if r["annual"] is not None else "—"
        print(f"{i:<4}{r['company']:<14}{r['role']:<12}{str(annual):<10}"
              f"{r['salary_score']:<8}{r['dev_score']:<8}{r['bal_score']:<8}{r['total']:<8}")
    top = results[0]
    gap = results[0]["total"] - results[1]["total"]
    print(f"\n🏆 综合最优: {top['company']}（总分 {top['total']}，最强项: {top['best']}，最弱项: {top['worst']}）")
    if gap < 0.3:
        print("分差 < 0.3，视为打平：建议用你最在意的单一因素做最终裁决。")
    elif gap < 0.8:
        print("分差 0.3-0.8，倾向第一名：但先复核低分者是否有可谈空间（如谈薪提升薪资分）。")
    else:
        print("分差 > 0.8，明确倾向第一名。")
    for r in results:
        print(f"  · {r['company']}: 最强 {r['best']} / 最弱 {r['worst']}")


def main():
    parser = argparse.ArgumentParser(description="求职投递跟踪表管理")
    sub = parser.add_subparsers(dest="command", required=True)

    p = sub.add_parser("init", help="初始化跟踪表")
    p.add_argument("file", nargs="?", default=DEFAULT_FILE)
    p.add_argument("--force", action="store_true", help="覆盖已存在的文件")
    p.set_defaults(func=cmd_init)

    p = sub.add_parser("add", help="添加投递记录")
    p.add_argument("file")
    p.add_argument("--company", required=True)
    p.add_argument("--role", default="")
    p.add_argument("--location", default="")
    p.add_argument("--channel", default="")
    p.add_argument("--date", default="")
    p.add_argument("--salary", default="")
    p.add_argument("--status", default="投递中")
    p.add_argument("--followup", default="")
    p.add_argument("--note", default="")
    p.set_defaults(func=cmd_add)

    p = sub.add_parser("update", help="更新投递状态/备注")
    p.add_argument("file")
    p.add_argument("--company", required=True)
    p.add_argument("--role", default=None)
    p.add_argument("--status", default=None)
    p.add_argument("--note", default=None)
    p.set_defaults(func=cmd_update)

    p = sub.add_parser("list", help="列出投递记录")
    p.add_argument("file")
    p.add_argument("--status", default=None)
    p.set_defaults(func=cmd_list)

    p = sub.add_parser("stats", help="投递统计概览")
    p.add_argument("file")
    p.set_defaults(func=cmd_stats)

    p = sub.add_parser("add-offer", help="录入 Offer 数据")
    p.add_argument("file")
    p.add_argument("--company", required=True)
    p.add_argument("--role", default="")
    p.add_argument("--month-salary", type=float, default=None, dest="month_salary")
    p.add_argument("--annual", type=float, default=None, help="年薪(万)")
    p.add_argument("--stock", type=float, default=None, help="股票期权(万/年)")
    p.add_argument("--bonus-months", type=float, default=None, dest="bonus_months", help="年终奖(月)")
    p.add_argument("--social-insurance", default="", dest="social_insurance", help="五险一金比例")
    p.add_argument("--scores", default=None, help="5个评分: 技术匹配,成长空间,稳定性,工作轻松度,通勤便利")
    p.add_argument("--note", default="")
    p.set_defaults(func=cmd_add_offer)

    p = sub.add_parser("compare", help="Offer 加权对比")
    p.add_argument("file")
    p.add_argument("--weights", default="0.4,0.3,0.3", help="薪资包,发展,平衡 权重")
    p.set_defaults(func=cmd_compare)

    args = parser.parse_args()
    try:
        args.func(args)
    except FileNotFoundError:
        print(f"文件不存在: {args.file}，请先运行 init 初始化。")
        sys.exit(1)
    except KeyError:
        print(f"文件缺少预期工作表，请确认使用 init 生成的跟踪表: {args.file}")
        sys.exit(1)


if __name__ == "__main__":
    main()
